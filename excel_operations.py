import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
import os

def read_excel_file(file_path, sheet_name=None):
    """
    Read data from an Excel file
    
    Args:
        file_path (str): Path to the Excel file
        sheet_name (str, optional): Name of the sheet to read. If None, reads the first sheet
    
    Returns:
        pandas.DataFrame: Data from the Excel file
    """
    try:
        # Read Excel file
        if sheet_name:
            df = pd.read_excel(file_path, sheet_name=sheet_name)
        else:
            df = pd.read_excel(file_path)
        
        print(f"Successfully read Excel file: {file_path}")
        print(f"Data shape: {df.shape}")
        print(f"Columns: {list(df.columns)}")
        
        return df
    
    except FileNotFoundError:
        print(f"Error: File '{file_path}' not found.")
        return None
    except Exception as e:
        print(f"Error reading Excel file: {str(e)}")
        return None

def write_excel_file(data, file_path, sheet_name='Sheet1', index=False):
    """
    Write data to an Excel file
    
    Args:
        data (pandas.DataFrame): Data to write
        file_path (str): Path where to save the Excel file
        sheet_name (str): Name of the sheet
        index (bool): Whether to include row indices
    """
    try:
        # Write to Excel file
        with pd.ExcelWriter(file_path, engine='openpyxl') as writer:
            data.to_excel(writer, sheet_name=sheet_name, index=index)
        
        print(f"Successfully wrote data to: {file_path}")
        
    except Exception as e:
        print(f"Error writing Excel file: {str(e)}")

def write_excel_with_formatting(data, file_path, sheet_name='Sheet1'):
    """
    Write data to Excel with custom formatting
    
    Args:
        data (pandas.DataFrame): Data to write
        file_path (str): Path where to save the Excel file
        sheet_name (str): Name of the sheet
    """
    try:
        # Create a workbook and select the active worksheet
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name
        
        # Write headers with formatting
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_alignment = Alignment(horizontal="center")
        
        # Write column headers
        for col_num, column_title in enumerate(data.columns, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.value = column_title
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
        
        # Write data
        for row_num, row_data in enumerate(data.values, 2):
            for col_num, cell_value in enumerate(row_data, 1):
                ws.cell(row=row_num, column=col_num, value=cell_value)
        
        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        # Save the workbook
        wb.save(file_path)
        print(f"Successfully wrote formatted Excel file: {file_path}")
        
    except Exception as e:
        print(f"Error writing formatted Excel file: {str(e)}")

def create_sample_data():
    """
    Create sample data for demonstration
    
    Returns:
        pandas.DataFrame: Sample data
    """
    sample_data = {
        'Name': ['Alice Johnson', 'Bob Smith', 'Charlie Brown', 'Diana Prince', 'Eva Garcia'],
        'Age': [25, 30, 35, 28, 32],
        'Department': ['Engineering', 'Marketing', 'Sales', 'HR', 'Finance'],
        'Salary': [75000, 65000, 55000, 60000, 70000],
        'Years_Experience': [3, 7, 10, 5, 8]
    }
    
    return pd.DataFrame(sample_data)

def read_multiple_sheets(file_path):
    """
    Read all sheets from an Excel file
    
    Args:
        file_path (str): Path to the Excel file
    
    Returns:
        dict: Dictionary with sheet names as keys and DataFrames as values
    """
    try:
        # Read all sheets
        all_sheets = pd.read_excel(file_path, sheet_name=None)
        
        print(f"Found {len(all_sheets)} sheets in {file_path}:")
        for sheet_name, df in all_sheets.items():
            print(f"  - {sheet_name}: {df.shape}")
        
        return all_sheets
    
    except Exception as e:
        print(f"Error reading multiple sheets: {str(e)}")
        return None

def write_multiple_sheets(data_dict, file_path):
    """
    Write multiple DataFrames to different sheets in one Excel file
    
    Args:
        data_dict (dict): Dictionary with sheet names as keys and DataFrames as values
        file_path (str): Path where to save the Excel file
    """
    try:
        with pd.ExcelWriter(file_path, engine='openpyxl') as writer:
            for sheet_name, df in data_dict.items():
                df.to_excel(writer, sheet_name=sheet_name, index=False)
        
        print(f"Successfully wrote {len(data_dict)} sheets to: {file_path}")
        
    except Exception as e:
        print(f"Error writing multiple sheets: {str(e)}")

def main():
    """
    Main function demonstrating Excel operations
    """
    print("=== Excel File Operations Demo ===\n")
    
    # 1. Create sample data
    print("1. Creating sample data...")
    sample_df = create_sample_data()
    print(sample_df)
    print()
    
    # 2. Write sample data to Excel
    print("2. Writing sample data to Excel...")
    output_file = "sample_data.xlsx"
    write_excel_file(sample_df, output_file)
    print()
    
    # 3. Read the Excel file back
    print("3. Reading Excel file...")
    read_df = read_excel_file(output_file)
    if read_df is not None:
        print("\nFirst 3 rows of read data:")
        print(read_df.head(3))
    print()
    
    # 4. Write with custom formatting
    print("4. Writing Excel file with formatting...")
    formatted_file = "formatted_data.xlsx"
    write_excel_with_formatting(sample_df, formatted_file)
    print()
    
    # 5. Demonstrate multiple sheets
    print("5. Creating Excel file with multiple sheets...")
    # Create different datasets for different sheets
    engineering_data = sample_df[sample_df['Department'] == 'Engineering']
    marketing_data = sample_df[sample_df['Department'] == 'Marketing']
    other_data = sample_df[~sample_df['Department'].isin(['Engineering', 'Marketing'])]
    
    multi_sheet_data = {
        'Engineering': engineering_data,
        'Marketing': marketing_data,
        'Other_Departments': other_data,
        'All_Data': sample_df
    }
    
    multi_sheet_file = "multi_sheet_data.xlsx"
    write_multiple_sheets(multi_sheet_data, multi_sheet_file)
    print()
    
    # 6. Read multiple sheets
    print("6. Reading multiple sheets...")
    all_sheets = read_multiple_sheets(multi_sheet_file)
    print()
    
    # 7. Data manipulation example
    print("7. Data manipulation example...")
    if read_df is not None:
        # Add a new column
        read_df['Salary_Category'] = read_df['Salary'].apply(
            lambda x: 'High' if x >= 70000 else 'Medium' if x >= 60000 else 'Low'
        )
        
        # Save the modified data
        modified_file = "modified_data.xlsx"
        write_excel_file(read_df, modified_file)
        print(f"Added 'Salary_Category' column and saved to {modified_file}")
    
    print("\n=== Demo completed! ===")
    print("Files created:")
    print("- sample_data.xlsx (basic Excel file)")
    print("- formatted_data.xlsx (with formatting)")
    print("- multi_sheet_data.xlsx (multiple sheets)")
    print("- modified_data.xlsx (with additional column)")

if __name__ == "__main__":
    main()